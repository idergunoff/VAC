"""Run the real legacy application in a separate process; never import it in tests.

record writes a candidate only. compare never updates approved fixtures.
"""
import argparse
import builtins
import hashlib
import importlib.metadata
import json
import os
from pathlib import Path
import platform
import runpy
import shutil
import sys
import tempfile
import traceback

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('mode', choices=['record', 'compare'])
parser.add_argument('--device', choices=['new', 'old'], required=True)
parser.add_argument('--output', type=Path, required=True)
parser.add_argument('--baseline', type=Path)
parser.add_argument('--case', default='full', choices=['full', 'single', 'cancel-dir', 'empty-dir', 'cancel-file', 'invalid', 'no-data', 'missing', 'unequal', 'after-dir', 'missing-noise', 'missing-channel'])
args = parser.parse_args()
if args.baseline:
    args.baseline = args.baseline.resolve()
args.output = args.output.resolve()
if args.output.exists():
    parser.error('output must be a new directory')
if args.mode == 'compare' and not args.baseline:
    parser.error('compare requires --baseline')
args.output.mkdir(parents=True)
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
os.environ['MPLBACKEND'] = 'Agg'
sys.dont_write_bytecode = True
import numpy as np
import pandas as pd
from PyQt5 import QtCore, QtWidgets
import matplotlib.pyplot as plt
import docx


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + '\n', encoding='utf-8')


def doc_summary(path):
    d = docx.Document(path)
    def table(t):
        return [[[c.text, [table(n) for n in c.tables]] for c in r.cells] for r in t.rows]
    return {'paragraphs': [p.text for p in d.paragraphs], 'tables': [table(t) for t in d.tables],
            'images': sorted(hashlib.sha256(r.target_part.blob).hexdigest() for r in d.part.rels.values() if 'image' in r.reltype)}


manifest = {'upstream': '6a99815036984389e51d8f0c934069af4332e278',
            'compatibility': 'ENV-01 int(auto_mean_win), see compatibility.md',
            'device': args.device, 'case': args.case, 'python': platform.python_version(),
            'platform': platform.platform(), 'dependencies': {d.metadata['Name']: d.version for d in importlib.metadata.distributions()},
            'source_sha256': {p: sha(ROOT / p) for p in ['VAC_python.py', 'VAC_dialog.py', 'responsive_layout.py', 'vac/reference.py', 'vac/resources/template.docx', 'vac/legacy.py']},
            'inputs': [], 'read_order': [], 'steps': [], 'qt_errors': []}
original_open = builtins.open
original_exec = QtWidgets.QApplication.exec_
status = 0

with tempfile.TemporaryDirectory(prefix='VAC эталоны с пробелами ') as tmp:
    scratch = Path(tmp)
    os.environ['MPLCONFIGDIR'] = str(scratch / 'mpl')
    sample = ROOT / 'пример' / ('новое оборудование' if args.device == 'new' else 'старое оборудование')
    target = scratch / sample.name
    shutil.copytree(sample, target)

    for p in sorted(target.rglob('*')):
        if p.is_file():
            manifest['inputs'].append({'path': str(p.relative_to(target)), 'sha256': sha(p), 'size': p.stat().st_size})

    def tracked_open(file, *a, **kw):
        try:
            p = Path(file)
            if p.is_relative_to(target) and p.suffix.lower() in ['.lvm', '.twf']:
                manifest['read_order'].append(str(p.relative_to(target)))
        except TypeError:
            pass
        return original_open(file, *a, **kw)

    builtins.open = tracked_open
    os.chdir(scratch)
    sys.path.insert(0, str(ROOT))

    def qt_error(kind, value, tb):
        manifest['qt_errors'].append({'type': kind.__name__, 'message': str(value)})
        traceback.print_exception(kind, value, tb)
    sys.excepthook = qt_error

    def event_loop(app):
        from vac import legacy
        scope = vars(legacy)
        ui = scope['ui']
        suffix = '' if args.device == 'new' else '_old'
        def call(name):
            scope[name]()
        def settings():
            result = {}
            for w in scope['MainWindow'].findChildren(QtWidgets.QWidget):
                name = w.objectName()
                if isinstance(w, (QtWidgets.QSpinBox, QtWidgets.QDoubleSpinBox)):
                    result[name] = w.value()
                elif isinstance(w, QtWidgets.QAbstractButton):
                    result[name] = {'checked': w.isChecked(), 'enabled': w.isEnabled(), 'text': w.text()}
                elif isinstance(w, QtWidgets.QComboBox):
                    result[name] = w.currentText()
                elif isinstance(w, QtWidgets.QLineEdit):
                    result[name] = w.text()
                elif isinstance(w, QtWidgets.QDateEdit):
                    result[name] = w.date().toString('yyyy-MM-dd')
            return result
        def snapshot(name):
            directory = args.output / name
            directory.mkdir()
            tables = {}
            for key, value in list(scope.items()):
                if isinstance(value, pd.DataFrame) and not value.empty:
                    table = {'columns': list(value.columns), 'index': value.index.tolist(), 'data': {}}
                    arrays = {}
                    for i, column in enumerate(value.columns):
                        series = value[column]
                        if pd.api.types.is_numeric_dtype(series):
                            arrays[str(i)] = series.to_numpy()
                            table['data'][str(i)] = {'numeric': True, 'dtype': str(series.dtype)}
                        else:
                            table['data'][str(i)] = [None if pd.isna(v) else str(v) for v in series]
                    np.savez_compressed(directory / (key + '.npz'), **arrays)
                    tables[key] = table
            write_json(directory / 'tables.json', tables)
            write_json(directory / 'settings.json', settings())
            write_json(directory / 'status.json', {'label_info': ui.label_info.text(), 'k_oldtonew': scope['k_oldtonew']})
            write_json(directory / 'intervals.json', {k: scope[k] for k in ['int_undefine_cement', 'int_0_cement', 'int_1_cement', 'int_2_cement']})
            classes = {}
            for key, value in list(scope.items()):
                if isinstance(value, pd.DataFrame):
                    for column in value.columns:
                        if 'qual' in str(column):
                            seq = value[column].tolist()
                            edges = [0] + [i for i in range(1, len(seq)) if seq[i] != seq[i-1]] + [len(seq)]
                            classes[key + ':' + str(column)] = [[a, b-1, int(seq[a])] for a, b in zip(edges, edges[1:]) if a < b]
            write_json(directory / 'classes.json', classes)
            plots = {}
            axes = []
            for f in plt.get_fignums():
                for ax in plt.figure(f).axes:
                    number = len(axes)
                    axes.append({'xlabel': ax.get_xlabel(), 'ylabel': ax.get_ylabel(), 'title': ax.get_title(),
                                 'xlim': list(ax.get_xlim()), 'ylim': list(ax.get_ylim()), 'lines': len(ax.lines)})
                    for j, line in enumerate(ax.lines):
                        plots[f'{number}:{j}:x'] = np.asarray(line.get_xdata(), dtype=float)
                        plots[f'{number}:{j}:y'] = np.asarray(line.get_ydata(), dtype=float)
            write_json(directory / 'matplotlib.json', axes)
            np.savez_compressed(directory / 'matplotlib.npz', **plots)
            plot = ui.graphicsView.getPlotItem()
            curves = {}
            for i, curve in enumerate(plot.listDataItems()):
                x, y = curve.getData()
                if x is not None:
                    curves[str(i) + 'x'] = x
                    curves[str(i) + 'y'] = y
            np.savez_compressed(directory / 'curves.npz', **curves)
            write_json(directory / 'axes.json', {k: {'text': v.labelText, 'units': v.labelUnits} for k, v in ((k, a['item']) for k, a in plot.axes.items())})
        def step(name, action):
            before = len(manifest['qt_errors'])
            try:
                action()
                snapshot(name)
                manifest['steps'].append({'name': name, 'status': 'ok' if len(manifest['qt_errors']) == before else 'qt-error'})
                print(name, manifest['steps'][-1]['status'], flush=True)
            except Exception as e:
                manifest['steps'].append({'name': name, 'status': 'error', 'type': type(e).__name__, 'message': str(e).replace(str(scratch), '<scratch>')})
                traceback.print_exc()
        def select(numbers):
            panel = ui.new_device if args.device == 'new' else ui.old_device
            for w in panel.findChildren(QtWidgets.QCheckBox):
                w.setChecked(int(w.text()[:2]) in numbers)
        def run():
            try:
                ui.dateEdit.setDate(QtCore.QDate(2026, 10, 19))
                ui.lineEdit_number_skv.setText('baseline')
                QtWidgets.QFileDialog.getExistingDirectory = staticmethod(lambda *a, **kw: str(target))
                load = 'open_dir' + suffix
                if args.case != 'full':
                    if args.case in ['cancel-dir', 'empty-dir', 'missing', 'missing-noise', 'missing-channel']:
                        if args.case in ['missing-noise', 'missing-channel']:
                            if args.device == 'new':
                                order = [Path(d) / f for d, _, files in os.walk(target) for f in sorted(files, reverse=True) if f.endswith('_5150m_c.lvm')]
                                remove = 8 if args.case == 'missing-noise' else 0
                            else:
                                order = [target / f for f in os.listdir(target) if f.endswith('.TWF')]
                                remove = 0 if args.case == 'missing-noise' else 1
                            removed = order[remove]
                            manifest['removed_input'] = str(removed.relative_to(target))
                            removed.unlink()
                            directory = str(target)
                        elif args.case == 'cancel-dir':
                            directory = ''
                        else:
                            empty = scratch / 'empty'; empty.mkdir()
                            if args.case == 'missing':
                                shutil.copy2(next(target.rglob('*.lvm' if args.device == 'new' else '*.TWF')), empty)
                            directory = str(empty)
                        QtWidgets.QFileDialog.getExistingDirectory = staticmethod(lambda *a, **kw: directory)
                        step(args.case, lambda: call(load))
                    elif args.case == 'no-data':
                        step(args.case, lambda: call('calc_to_int' + suffix))
                    else:
                        if args.case == 'after-dir':
                            step('load-directory', lambda: call(load))
                        path = next(target.rglob('*.lvm' if args.device == 'new' else '*.TWF'))
                        if args.case == 'invalid':
                            path = scratch / 'invalid'; path.write_bytes(b'not a measurement')
                        QtWidgets.QFileDialog.getOpenFileName = staticmethod(lambda *a, **kw: ('' if args.case == 'cancel-file' else str(path), ''))
                        step(args.case, lambda: call('open_sig' + suffix))
                        if args.case == 'unequal':
                            if args.device == 'new':
                                lines = path.read_text().splitlines()
                                path = target / 'short.lvm'
                                path.write_text('\n'.join(lines[:len(lines)//2]) + '\n')
                            else:
                                data = path.read_bytes()
                                path = target / 'short.TWF'
                                path.write_bytes(data[:130 + ((len(data)-130)//8)*4])
                            QtWidgets.QFileDialog.getOpenFileName = staticmethod(lambda *a, **kw: (str(path), ''))
                            step('short-file', lambda: call('open_sig' + suffix))
                            step('select', lambda: select([1, 2]))
                            step('sum-unequal', lambda: call('sum1' + suffix))
                    return
                step('01-load', lambda: call(load))
                step('02-select', lambda: select([1, 2]))
                def interval():
                    ui.doubleSpinBox_int_min.setValue(10)
                    ui.doubleSpinBox_int_max.setValue(150)
                    call('calc')
                step('03-interval', interval)
                def mean():
                    if args.device == 'new':
                        ui.spinBox_mean.setValue(60); call('calc_mean')
                    getattr(ui, 'spinBox_mean_win' + suffix).setValue(40)
                    getattr(ui, 'doubleSpinBox_coeff_norm' + suffix).setValue(0.2)
                step('04-mean-norm', mean)
                step('05-sum', lambda: call('sum1' + suffix))
                step('06-quality', lambda: call('recalc_qual_' + args.device))
                step('07-cement-source', lambda: (select([1]), call('calc_int_cement')))
                def correct():
                    ui.checkBox_shift_corr.setChecked(True)
                    ui.doubleSpinBox_x1_line.setValue(20)
                    ui.doubleSpinBox_x2_line.setValue(30)
                    call('corr_sig'); call('fix_corr')
                    ui.radioButton_NC.setChecked(True)
                    ui.doubleSpinBox_undefine_min.setValue(40)
                    ui.doubleSpinBox_undefine_max.setValue(45)
                    call('add_undefine')
                step('08-correction', correct)
                step('09-cementogram', lambda: call('cementogramma'))
                step('10-save', lambda: call('save_cement'))
                for p in scratch.glob('*.docx'):
                    if p.name != 'template.docx':
                        shutil.copy2(p, args.output / 'report.docx')
                        write_json(args.output / 'report.json', doc_summary(p))
                for p in scratch.glob('*.vac'):
                    shutil.copy2(p, args.output / 'settings.vac')
                    QtWidgets.QFileDialog.getOpenFileName = staticmethod(lambda *a, **kw: (str(p), ''))
                    step('11-restore', lambda: call('open_options'))
                step('12-reload', lambda: call(load))
            finally:
                app.quit()
        QtCore.QTimer.singleShot(0, run)
        original_exec()
        return 0
    QtWidgets.QApplication.exec_ = event_loop
    try:
        runpy.run_path(str(ROOT / 'VAC_python.py'), run_name='__main__')
    except SystemExit as e:
        if e.code:
            raise
    finally:
        builtins.open = original_open
        write_json(args.output / 'manifest.json', manifest)
        plt.close('all')


os.chdir(ROOT)
from regression_compare import compare

if args.mode == 'compare':
    compare(args.baseline.resolve(), args.output)
else:
    print('Candidate recorded:', args.output)
