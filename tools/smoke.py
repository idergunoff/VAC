"""Headless launch/sample probe; no numerical or visual acceptance claim."""
import argparse
import json
import os
from pathlib import Path
import runpy
import shutil
import sys
import tempfile
import traceback

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument('--scenario', choices=['launch', 'new', 'old'], default='launch')
args = parser.parse_args()
root = Path(__file__).resolve().parents[1]
os.environ['QT_QPA_PLATFORM'] = 'offscreen'
os.environ['MPLBACKEND'] = 'Agg'
sys.dont_write_bytecode = True

with tempfile.TemporaryDirectory(prefix='VAC проверка с пробелами ') as scratch:
    os.environ['MPLCONFIGDIR'] = str(Path(scratch) / 'matplotlib')
    from PyQt5 import QtCore, QtWidgets
    import docx
    result = {'scenario': args.scenario, 'passed': False, 'exceptions': []}
    original_exec = QtWidgets.QApplication.exec_

    def record_error(kind, value, tb):
        result['exceptions'].append(''.join(traceback.format_exception(kind, value, tb)))
        traceback.print_exception(kind, value, tb)

    sys.excepthook = record_error

    def event_loop(app):
        from vac import legacy
        scope = vars(legacy)

        def check():
            try:
                assert scope['MainWindow'].isVisible(), 'Window not visible'
                assert scope['ui'].graphicsView is not None
                template = docx.Document(str(template_path()))
                assert len(template.tables) >= 2
                result['template_tables'] = len(template.tables)
                if args.scenario != 'launch':
                    sample = root / 'пример' / ('новое оборудование' if args.scenario == 'new' else 'старое оборудование')
                    # Copy data to exercise spaces and Cyrillic in actual input paths.
                    target = Path(scratch) / sample.name
                    shutil.copytree(sample, target)
                    QtWidgets.QFileDialog.getExistingDirectory = staticmethod(lambda *a, **kw: str(target))
                    scope['open_dir' if args.scenario == 'new' else 'open_dir_old']()
                    tab = scope['all_signals' if args.scenario == 'new' else 'all_signals_old']
                    assert not tab.empty
                    result['shape'] = list(tab.shape)
                result['passed'] = not result['exceptions']
            except Exception:
                record_error(*sys.exc_info())
            finally:
                result['outputs'] = sorted(p.name for p in Path(scratch).iterdir() if p.is_file())
                QtWidgets.QApplication.instance().quit()

        QtCore.QTimer.singleShot(0, check)
        original_exec()
        print(json.dumps(result, ensure_ascii=False, indent=2), flush=True)
        return int(not result['passed'])

    QtWidgets.QApplication.exec_ = event_loop
    sys.path.insert(0, str(root))
    from vac.resources import template_path
    os.chdir(scratch)
    runpy.run_path(str(root / 'VAC_python.py'), run_name='__main__')
