"""Word template population, independent of Qt and controller state."""
import docx
from pathlib import Path
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from copy import deepcopy


def build_report(table_data, frequency, values, template, image_path='temp.png'):
    doc = docx.Document(template)
    table = doc.tables[0]
    table2 = doc.tables[1]
    table.rows[0].cells[0].paragraphs[0].runs[3].text = values['lineEdit_number_skv.text']
    doc.paragraphs[4].runs[2].text = values['lineEdit_number_skv.text']
    doc.paragraphs[9].runs[5].text = values['lineEdit_number_skv.text']
    table.rows[1].cells[0].paragraphs[0].runs[0].text = values['comboBox_pipe_type.currentText']
    if values['comboBox_pipe_type.currentText'] == 'Кондуктор':
        tube = 'кондуктора'
    elif values['comboBox_pipe_type.currentText'] == 'Колонна':
        tube = 'колонны'
    elif values['comboBox_pipe_type.currentText'] == 'Тех. колонна':
        tube = 'тех. колонны'
    elif values['comboBox_pipe_type.currentText'] == 'Направление':
        tube = 'направления'
    else:
        tube = 'кондуктора'
    doc.paragraphs[9].runs[3].text = tube
    doc.paragraphs[11].runs[1].text = tube
    table.rows[1].cells[0].paragraphs[0].runs[3].text = values['lineEdit_pipe_diameter.text']
    doc.paragraphs[9].runs[7].text = values['lineEdit_pipe_diameter.text']
    table.rows[2].cells[0].paragraphs[0].runs[1].text = values['lineEdit_client.text']
    doc.paragraphs[6].runs[1].text = values['lineEdit_client.text']
    table.rows[3].cells[0].paragraphs[0].clear()
    doc.paragraphs[7].clear()
    client2 = table.rows[3].cells[0].paragraphs[0].add_run(values['lineEdit_client2.text'])
    client2.font.size = Pt(11)
    client2.font.name = 'Arial'
    client2 = doc.paragraphs[7].add_run(values['lineEdit_client2.text'])
    client2.font.size = Pt(12)
    client2.font.name = 'Times New Roman'
    table.rows[1].cells[1].paragraphs[0].runs[1].text = str(round(table_data['depth'].min(), 1))
    table.rows[1].cells[1].paragraphs[0].runs[5].text = str(float(values['doubleSpinBox_int_max.value']))
    doc.paragraphs[9].runs[10].text = str(round(table_data['depth'].min(), 1))
    doc.paragraphs[9].runs[12].text = str(float(values['doubleSpinBox_int_max.value']))
    int_study = table_data['depth'].max() - table_data['depth'].min()
    table.rows[2].cells[1].paragraphs[0].runs[1].text = values['dateEdit.text']
    table.rows[9].cells[2].paragraphs[0].clear()
    interpreter = table.rows[9].cells[2].paragraphs[0].add_run(values['lineEdit_interpreter.text'])
    interpreter.font.size = Pt(11)
    interpreter.font.name = 'Times New Roman'
    table2.rows[1].cells[1].paragraphs[0].clear()
    interpreter = table2.rows[1].cells[1].paragraphs[0].add_run(values['lineEdit_interpreter.text'])
    interpreter.font.size = Pt(11)
    interpreter.font.name = 'Times New Roman'
    table.rows[6].cells[0]._element.clear_content()
    table.rows[6].cells[0].add_paragraph().add_run().add_picture(str(image_path))
    table.rows[6].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    table.rows[7].cells[2].paragraphs[0].runs[5].text = frequency
    n_row = 1
    h_def = 0
    if values['checkBox_golfstrim.checkState'] == 2:
        if table_data['quality'][0] == 0:
            table2.rows[0].cells[0].tables[0].add_row()
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].add_run(str(round(table_data['depth'][0], 1)))
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, quality in enumerate(table_data['quality']):
            if 0 < i < len(table_data['quality']) - 1:
                if quality == 0:
                    if table_data['quality'][i - 1] != 0:
                        table2.rows[0].cells[0].tables[0].add_row()
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].add_run(str(round(table_data['depth'][i], 1)))
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if table_data['quality'][i + 1] != 0:
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].add_run(str(round(table_data['depth'][i], 1)))
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        h = float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].text) - float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].text)
                        h_def = h_def + h
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].add_run(str(round(h, 1)))
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].add_run('Качественный цемент')
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        n_row += 1
        if table_data['quality'][len(table_data['quality']) - 1] == 0:
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].add_run(str(round(table_data['depth'][len(table_data['quality']) - 1], 1)))
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            h = float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].text) - float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].text)
            h_def = h_def + h
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].add_run(str(round(h, 1)))
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].add_run('Качественный цемент')
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        if table_data['quality'][0] > 0:
            table2.rows[0].cells[0].tables[0].add_row()
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].add_run(str(round(table_data['depth'][0], 1)))
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for i, quality in enumerate(table_data['quality']):
            if 0 < i < len(table_data['quality']) - 1:
                if quality > 0:
                    if table_data['quality'][i - 1] != quality:
                        table2.rows[0].cells[0].tables[0].add_row()
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].add_run(str(round(table_data['depth'][i], 1)))
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    if table_data['quality'][i + 1] != quality:
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].add_run(str(round(table_data['depth'][i], 1)))
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        h = float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].text) - float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].text)
                        h_def = h_def + h
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].add_run(str(round(h, 1)))
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        if quality == 1:
                            defect = 'частичный цемент'
                        elif quality == 2:
                            defect = 'отсутствует цемент'
                        elif quality == 3:
                            defect = 'неопределённый цемент'
                        note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].add_run(defect)
                        note.font.size = Pt(11)
                        note.font.name = 'Times New Roman'
                        table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                        n_row += 1
        if table_data['quality'][len(table_data['quality']) - 1] > 0:
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].add_run(str(round(table_data['depth'][len(table_data['quality']) - 1], 1)))
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            h = float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[1].paragraphs[0].text) - float(table2.rows[0].cells[0].tables[0].rows[n_row].cells[0].paragraphs[0].text)
            h_def = h_def + h
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].add_run(str(round(h, 1)))
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if table_data['quality'][len(table_data['quality']) - 1] == 1:
                defect = 'частичный цемент'
            elif table_data['quality'][len(table_data['quality']) - 1] == 2:
                defect = 'отсутствует цемент'
            elif table_data['quality'][len(table_data['quality']) - 1] == 3:
                defect = 'неопределённый цемент'
            note = table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].add_run(defect)
            note.font.size = Pt(11)
            note.font.name = 'Times New Roman'
            table2.rows[0].cells[0].tables[0].rows[n_row].cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    if values['checkBox_golfstrim.checkState'] == 2:
        k_def = h_def / int_study
    else:
        k_def = 1 - h_def / int_study
    table.rows[7].cells[0].paragraphs[0].runs[7].text = str(round(k_def, 1))
    if table_data['depth'].max() > 400:
        new_tbl = deepcopy(table._tbl)
        table._tbl.addnext(new_tbl)
        table3 = doc.tables[1]
        table3.rows[6].cells[0]._element.clear_content()
        page = Path(image_path)
        table3.rows[6].cells[0].add_paragraph().add_run().add_picture(str(page.with_name(page.stem + '1' + page.suffix)))
        table3.rows[6].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        table3.rows[0].cells[3].paragraphs[0].runs[0].text = 'лист 2'
    if table_data['depth'].max() > 800:
        new_tbl2 = deepcopy(table3._tbl)
        table3._tbl.addnext(new_tbl2)
        table4 = doc.tables[2]
        table4.rows[6].cells[0]._element.clear_content()
        page = Path(image_path)
        table4.rows[6].cells[0].add_paragraph().add_run().add_picture(str(page.with_name(page.stem + '2' + page.suffix)))
        table4.rows[6].cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        table4.rows[0].cells[3].paragraphs[0].runs[0].text = 'лист 3'
    return doc
